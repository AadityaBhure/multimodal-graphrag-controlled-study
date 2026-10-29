"""Build the B.Tech PBL project report (Word) in the format of the department template (assets/template.docx).

Layout taken from the template: A4; margins 1.3 in left, 1.0 in right, top and bottom; Times New Roman 12 pt,
justified; SIT Nagpur cover page and certificate with the institute logo; declaration, acknowledgement, abstract;
table of contents, list of tables, list of figures and list of abbreviations as tables with page numbers; front
matter numbered i, ii, ...; chapters numbered from 1, each starting on a new page with "Chapter n" and its title
centred in 16 pt bold; sections "n.m" in bold; references in IEEE format; appendices.

Page numbers in the contents and lists are PAGEREF fields to bookmarks on every heading and caption, filled in when
../../paper/ieee/build/finalize_ieee.py updates the fields; equations are built into Word equations by the same step.

Markup (report.md):
  %key value                       metadata (title, title_cover, student, prn, guide, designation, ay, ...)
  @acknowledgement / @abstract ... @end, @keywords ...
  @abbrev ... @end                 one "ABBR | expansion" per line
  # Title {#ch:x}                  chapter          #A Title {#app:x}   appendix          # *Title   unnumbered (References)
  ## Title {#sec:x}                section n.m      ### Title           subsection n.m.k  #### Title   bold run-in heading
  @fig label | file | caption | width-in          @table label | caption | widths   (pipe table follows)
  @results label | RQx | metrics | caption [| methods]     @eq label | UnicodeMath     @alg label | title ... @end
  @code ... @end                   monospaced listing    @refs   reference list       @pagebreak
Inline: [@key;key] IEEE citations, [[fig:x]] [[tab:x]] [[eq:x]] [[alg:x]] [[sec:x]] [[ch:x]] [[app:x]], {{KEY}} frozen
results, $...$ inline math, *italic*, **bold**, `code`.
"""
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
sys.path.insert(0, str(ROOT / 'documentation/paper/ieee/build'))
import build_ieee as W                                   # noqa: E402  table helpers, row splitter, quotes, number binding
from refs import REFS                                    # noqa: E402  IEEE-format references

BD = W.BD
OUT = HERE.parent / 'Multi-Modal_GraphRAG_Project_Report.docx'
TEMPLATE = HERE / 'assets' / 'template.docx'
LOGO = HERE / 'assets' / 'sit_logo.png'
FIG_DIRS = [HERE.parent / 'figures', ROOT / 'documentation/paper/scirep/figures', ROOT / 'documentation/paper/ieee/figures',
            ROOT / 'documentation/project_documentation']
FONT, MONO = 'Times New Roman', 'Consolas'
BODY = 12
LINE = 1.5                                               # line spacing of body text
TEXTW = 8.27 - 1.3 - 1.0                                 # 5.97 in
MATH_L, MATH_R = W.MATH_L, W.MATH_R
P6 = json.loads((ROOT / 'documentation/phase6_closeout/comparison.json').read_text(encoding='utf-8'))


def bound(key):
    if key.startswith('P6.'):
        _, b, metric = key.split('.', 2)
        row = next(r for r in P6 if r['baseline'] == b and r['cohort'] == 'overall' and r['metric'] == metric)
        return BD.fnum(row['mean'], 3)
    return W.bound(key)


# =============================================================================== low-level helpers
def font(run, size=BODY, bold=None, italic=None, face=FONT, color=None):
    run.font.name = face
    rf = run._r.get_or_add_rPr().get_or_add_rFonts()
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), face)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


TOKENS = re.compile(r'(⟪[^⟫]*⟫|\*\*.+?\*\*|`[^`]+`|(?<![\w*])\*[^*\s][^*]*?\*(?![\w*]))')
SYMBOLS = set('★✓✗')


def add(p, text, size=BODY, bold=False, italic=False, face=FONT, color=None):
    for tok in TOKENS.split(text):
        if not tok:
            continue
        if tok[0] == MATH_L:
            r = p.add_run(tok)
            font(r, size, bold, italic, face, color)
        elif tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            add(p, tok[2:-2], size, True, italic, face, color)
        elif tok.startswith('`') and tok.endswith('`') and len(tok) > 2:
            r = p.add_run(tok[1:-1])
            font(r, size - 1.5, bold, False, MONO, color)
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            add(p, tok[1:-1], size, bold, True, face, color)
        else:
            for part in re.split('([%s])' % ''.join(SYMBOLS), tok):
                if part:
                    r = p.add_run(part)
                    font(r, size, bold, italic, 'Segoe UI Symbol' if part in SYMBOLS else face, color)
    return p


def fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=6, first=0.0, left=0.0, keep=False, line=LINE, page_break=False):
    f = p.paragraph_format
    p.alignment = align
    f.space_before, f.space_after = Pt(before), Pt(after)
    f.first_line_indent = Inches(first) if first else None
    f.left_indent = Inches(left) if left else None
    f.keep_with_next = keep
    f.widow_control = True
    f.page_break_before = page_break
    f.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    f.line_spacing = line
    return p


def field(p, instr, placeholder='1', size=BODY, bold=False):
    def run():
        r = p.add_run()
        font(r, size, bold)
        return r
    r = run()
    el = OxmlElement('w:fldChar')
    el.set(qn('w:fldCharType'), 'begin')
    r._r.append(el)
    r = run()
    el = OxmlElement('w:instrText')
    el.set(qn('xml:space'), 'preserve')
    el.text = ' %s ' % instr
    r._r.append(el)
    r = run()
    el = OxmlElement('w:fldChar')
    el.set(qn('w:fldCharType'), 'separate')
    r._r.append(el)
    r = run()
    r.text = placeholder
    r = run()
    el = OxmlElement('w:fldChar')
    el.set(qn('w:fldCharType'), 'end')
    r._r.append(el)


BM = [0]


def bookmark(p, name):
    BM[0] += 1
    s = OxmlElement('w:bookmarkStart')
    s.set(qn('w:id'), str(BM[0]))
    s.set(qn('w:name'), name)
    e = OxmlElement('w:bookmarkEnd')
    e.set(qn('w:id'), str(BM[0]))
    p._p.insert(1 if p._p.pPr is not None else 0, s)
    p._p.append(e)


def shade(cell, fill):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def find_fig(name):
    for d in FIG_DIRS:
        if (d / name).exists():
            return d / name
    raise FileNotFoundError(name)


def short(caption):
    """List-of-figures text: the caption up to its first sentence end."""
    c = re.sub(r'\s+', ' ', caption.strip())
    m = re.search(r'(?<![A-Z][a-z])(?<!\b[A-Z])(?<!e\.g)(?<!i\.e)(?<!vs)(?<!Fig)\.\s', c)
    return c[:m.start()] if m and m.start() > 25 else c.rstrip('.')


# =============================================================================== labels
class Labels:
    def __init__(self):
        self.map = {}                                     # 'fig:x' -> '4.2'
        self.cites = []

    def ref(self, key):
        kind, _, lab = key.partition(':')
        v = self.map[key]
        return {'fig': 'Figure %s', 'tab': 'Table %s', 'eq': 'Equation (%s)', 'alg': 'Algorithm %s', 'sec': 'Section %s',
                'ch': 'Chapter %s', 'app': 'Appendix %s'}[kind] % v

    def cite(self, keys):
        nums = []
        for k in (k.strip() for k in keys):
            if k not in REFS:
                raise KeyError('unknown reference ' + k)
            if k not in self.cites:
                self.cites.append(k)
            nums.append(self.cites.index(k) + 1)
        nums = sorted(set(nums))
        parts, i = [], 0
        while i < len(nums):
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                j += 1
            parts.append('[%d]–[%d]' % (nums[i], nums[j]) if j - i >= 2 else ', '.join('[%d]' % n for n in nums[i:j + 1]))
            i = j + 1
        return ', '.join(parts)


def expand(text, L):
    text = re.sub(r'\{\{([^}]+)\}\}', lambda m: bound(m.group(1).strip()), text)
    text = re.sub(r'\[\[((?:fig|tab|eq|alg|sec|ch|app):[\w\-]+)\]\]', lambda m: L.ref(m.group(1)), text)
    text = re.sub(r'\[@([^\]]+)\]', lambda m: L.cite(m.group(1).split(';')), text)
    return W.smart(W.math_inline(text))


def prescan(lines, L):
    """Number chapters, sections, figures, tables, equations and algorithms before rendering (forward references)."""
    ch, app, sec, sub = 0, 0, 0, 0
    cnt = {}
    for line in lines:
        s = line.strip()
        h = re.match(r'^(#{1,4}|#A)\s+(\*?)(.*?)\s*(\{#((?:ch|sec|app):[\w\-]+)\})?$', s)
        if h and not s.startswith('####'):
            mark, star, label = h.group(1), h.group(2), h.group(5)
            if mark == '#A':
                app += 1
                cur = chr(64 + app)
                sec = 0
                cnt = {}
                if label:
                    L.map[label] = cur
                chap = cur
            elif mark == '#':
                if not star:
                    ch += 1
                    sec = 0
                    cnt = {}
                    if label:
                        L.map[label] = str(ch)
                chap = str(ch)
            elif mark == '##':
                sec += 1
                sub = 0
                if label:
                    L.map[label] = '%s.%d' % (chap, sec)
            continue
        m = re.match(r'^@(fig|table|results|eq|alg)\s+([\w\-]+)\s*\|', s)
        if m:
            kind = {'fig': 'fig', 'table': 'tab', 'results': 'tab', 'eq': 'eq', 'alg': 'alg'}[m.group(1)]
            cnt[kind] = cnt.get(kind, 0) + 1
            L.map['%s:%s' % (kind, m.group(2))] = '%s.%d' % (chap, cnt[kind])


# =============================================================================== the report
class Report:
    def __init__(self):
        self.doc = Document(str(TEMPLATE))
        body = self.doc.element.body
        for el in list(body):
            if el.tag != qn('w:sectPr'):
                body.remove(el)
        self.L = Labels()
        self.toc, self.lot, self.lof = [], [], []
        self.chap = None
        self.sec = self.sub = 0
        self.cnt = {}
        self.first = True
        self.mark = 0
        self.placeholders = {}
        self._styles()
        s = self.doc.sections[0]
        s.page_width, s.page_height = Inches(8.27), Inches(11.69)
        s.left_margin, s.right_margin = Inches(1.3), Inches(1.0)
        s.top_margin, s.bottom_margin = Inches(1.0), Inches(1.0)
        s.footer_distance = Inches(0.43)

    def _styles(self):
        st = self.doc.styles
        n = st['Normal']
        n.font.name = FONT
        n.font.size = Pt(BODY)
        n.paragraph_format.space_before = Pt(0)
        n.paragraph_format.line_spacing = 1.0
        try:
            st['IEEE Equation']
        except KeyError:
            eq = st.add_style('IEEE Equation', WD_STYLE_TYPE.PARAGRAPH)
            eq.base_style = n
            eq.paragraph_format.space_before = Pt(6)
            eq.paragraph_format.space_after = Pt(6)
            eq.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            eq.font.size = Pt(BODY)
        for name in ('Heading 1', 'Heading 2', 'Heading 3'):
            h = st[name]
            h.font.name = FONT
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.font.small_caps = False
            h.font.all_caps = False
            h.font.italic = False
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.page_break_before = False
            h.paragraph_format.left_indent = Inches(0)
            h.paragraph_format.first_line_indent = Inches(0)

    def newmark(self):
        self.mark += 1
        return '_Ref_rp%04d' % self.mark

    # ------------------------------------------------------------------ sections and footers
    def _footer(self, section, fmt_='decimal', start=1, numbered=True):
        section.footer.is_linked_to_previous = False
        ft = section.footer
        p = ft.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        if numbered:
            fmt(p, WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.0)
            field(p, 'PAGE', '1', 11)
        sp = section._sectPr
        for old in sp.findall(qn('w:pgNumType')):
            sp.remove(old)
        pg = OxmlElement('w:pgNumType')
        pg.set(qn('w:fmt'), fmt_)
        pg.set(qn('w:start'), str(start))
        sp.append(pg)
        for t in sp.findall(qn('w:titlePg')):
            sp.remove(t)

    def new_section(self, fmt_, start):
        s = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._footer(s, fmt_, start)
        return s

    # ------------------------------------------------------------------ front matter
    def para(self, text='', size=BODY, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0, line=1.0):
        p = self.doc.add_paragraph()
        fmt(p, align, before=before, after=after, line=line)
        if text:
            add(p, text, size, bold, italic)
        return p

    def cover(self, M):
        p = self.para(after=14)
        p.add_run().add_picture(str(LOGO), width=Inches(5.9))
        self.para('A PROJECT REPORT', 12, True, after=6)
        self.para('ON', 12, True, after=6)
        self.para('“%s”' % M['title_cover'], 16, True, before=4, after=22, line=1.15)
        self.para('A project report submitted in partial fulfillment of the requirements for the degree of Bachelor of Technology '
                  'in Computer Science and Engineering', 12, after=26, line=1.15)
        self.para('BACHELOR OF TECHNOLOGY COMPUTER SCIENCE AND ENGINEERING', 14, after=26, line=1.15)
        self.para('Submitted By', 14, after=2)
        self.para('%s (PRN: %s)' % (M['student'], M['prn']), 14, True, after=26)
        self.para('UNDER THE GUIDANCE OF', 14, after=8)
        self.para(M['guide'], 14, True, after=6)
        self.para(M['designation'], 14, after=40)
        self.para('DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING', 12, True, after=2)
        self.para('AY %s' % M['ay'], 12, True)
        s = self.doc.sections[0]
        self._footer(s, numbered=False)

    def front_heading(self, title, toc=True):
        p = self.para(title, 14, True, after=18)
        self.take_break(p)
        if toc:
            bm = self.newmark()
            bookmark(p, bm)
            self.toc.append(('front', '', title.title().replace('Of', 'of'), bm))
        return p

    def certificate(self, M):
        self.new_section('lowerRoman', 1)
        p = self.para(after=10)
        p.add_run().add_picture(str(LOGO), width=Inches(5.9))
        self.para('DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING', 12, True, after=40)
        p = self.para('CERTIFICATE', 14, True, after=18)
        bm = self.newmark()
        bookmark(p, bm)
        self.toc.append(('front', '', 'Certificate', bm))
        p = self.doc.add_paragraph()
        fmt(p, before=0, after=60, line=LINE)
        add(p, 'This is to certify that the Project work entitled “**%s**” is carried out by **%s**, in partial fulfillment '
               'for the award of the degree of **Bachelor of Technology** in **Computer Science and Engineering**, Symbiosis '
               'International (Deemed University), Pune during the academic year %s.' % (M['title'], M['student'], M['ay_long']))
        t = self.doc.add_table(rows=3, cols=3)
        W.table_fixed(t, [2.6, 0.77, 2.6])
        names = [('Dr. Giridhar Urkude', 'PBL Guide'), ('Dr. Giridhar Urkude', 'PBL Coordinator'),
                 ('Dr. Sagarkumar S. Badhiye', 'Head of Department, CSE'), ('Dr. Nitin Rakesh', 'Director, SIT Nagpur')]
        if M.get('cert_guide'):
            names[0] = (M['cert_guide'], 'PBL Guide')
        for (r, c), (nm, role) in zip(((0, 0), (0, 2), (2, 0), (2, 2)), names):
            cell = t.cell(r, c)
            q = cell.paragraphs[0]
            fmt(q, WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.0)
            add(q, nm, 12, True)
            q = cell.add_paragraph()
            fmt(q, WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.0)
            add(q, role, 12)
        sp = t.rows[1].cells[0].paragraphs[0]
        fmt(sp, after=0, line=1.0)
        sp.paragraph_format.space_before = Pt(80)

    def declaration(self, M):
        self.page_break()
        self.front_heading('DECLARATION')
        p = self.doc.add_paragraph()
        fmt(p, after=14, line=LINE)
        add(p, 'I hereby declare that the project titled “**%s**” submitted to Symbiosis Institute of Technology, '
               'Constituent of Symbiosis International (Deemed University) Pune for the award of the degree of Bachelor of '
               'Technology in Computer Science and Engineering is a result of original research carried out by me. I understand '
               'that my report may be made electronically available to the public. It is further declared that the project report '
               'or any part thereof has not been previously submitted to any University or Institute for the award of any degree '
               'or diploma.' % M['title'])
        rows = [('Name of the Student', M['student']), ('PRN', M['prn']), ('Degree', 'Bachelor of Technology in CSE'),
                ('Department', 'Computer Science and Engineering'), ('Title of the project', M['title'])]
        for k, v in rows:
            q = self.doc.add_paragraph()
            fmt(q, WD_ALIGN_PARAGRAPH.LEFT, after=8, left=1.97, first=-1.97, line=1.15)
            q.paragraph_format.tab_stops.add_tab_stop(Inches(1.97))
            add(q, '%s\t: %s' % (k, v))
        for t, a, b in (('_________________________________', 70, 2), ('(Signature of the Student)', 0, 14), ('Date:', 0, 0)):
            q = self.doc.add_paragraph()
            fmt(q, WD_ALIGN_PARAGRAPH.LEFT, before=a, after=b, line=1.0)
            add(q, t)

    def page_break(self):
        self.pending_break = True

    def take_break(self, p):
        if getattr(self, 'pending_break', False):
            p.paragraph_format.page_break_before = True
            self.pending_break = False

    def text_page(self, title, paras):
        self.page_break()
        self.front_heading(title)
        for t in (x.strip() for x in paras):
            p = self.doc.add_paragraph()
            fmt(p, after=8, line=LINE)
            add(p, expand(t, self.L))

    def abstract(self, body, keywords):
        self.page_break()
        self.front_heading('ABSTRACT')
        for t in (x.strip() for x in body):
            p = self.doc.add_paragraph()
            fmt(p, after=8, line=LINE)
            add(p, expand(t, self.L))
        p = self.doc.add_paragraph()
        fmt(p, WD_ALIGN_PARAGRAPH.JUSTIFY, before=6, after=0, line=LINE)
        add(p, 'Keywords: ', bold=True)
        add(p, keywords)

    def lists(self, abbrevs):
        for key, title in (('toc', 'TABLE OF CONTENTS'), ('lot', 'LIST OF TABLES'), ('lof', 'LIST OF FIGURES')):
            self.page_break()
            self.front_heading(title)
            ph = self.doc.add_paragraph()
            fmt(ph, after=0, line=1.0)
            self.placeholders[key] = ph
        self.page_break()
        self.front_heading('LIST OF ABBREVIATIONS')
        t = self.doc.add_table(rows=1 + len(abbrevs), cols=2)
        W.table_fixed(t, [1.5, 4.3])
        for ri, (a, b) in enumerate([('Abbreviation', 'Full form')] + abbrevs):
            for ci, v in enumerate((a, b)):
                c = t.cell(ri, ci)
                W.cell_margins(c, t=30, b=30, l=80, r=80)
                W.cell_borders(c, **({'bottom': 6, 'top': 6} if ri == 0 else {}))
                q = c.paragraphs[0]
                fmt(q, WD_ALIGN_PARAGRAPH.LEFT, after=0, line=1.15)
                add(q, v, 12, bold=ri == 0 or ci == 0 and False)

    def fill_lists(self):
        def table_at(ph, rows, widths, header=True):
            t = self.doc.add_table(rows=len(rows) + (1 if header else 0), cols=3)
            W.table_fixed(t, widths)
            r0 = 0
            if header:
                c = t.cell(0, 2)
                q = c.paragraphs[0]
                fmt(q, WD_ALIGN_PARAGRAPH.RIGHT, after=4, line=1.0)
                add(q, 'Page', 12, True)
                r0 = 1
            for i, (kind, num, title, bm) in enumerate(rows):
                tr = t.rows[r0 + i]
                W.row_props(tr)
                bold = kind in ('front', 'chapter', 'end')
                if kind in ('front', 'chapter', 'end'):
                    a = tr.cells[0].merge(tr.cells[1])
                    q = a.paragraphs[0]
                    fmt(q, WD_ALIGN_PARAGRAPH.LEFT, before=6 if kind != 'front' else 2, after=2, line=1.0)
                    add(q, title, 12, bold)
                else:
                    q = tr.cells[0].paragraphs[0]
                    fmt(q, WD_ALIGN_PARAGRAPH.LEFT, before=2, after=2, line=1.0, left=0.25 if kind == 'sub' else 0)
                    add(q, num, 12 if kind != 'sub' else 11)
                    q = tr.cells[1].paragraphs[0]
                    fmt(q, WD_ALIGN_PARAGRAPH.LEFT, before=2, after=2, line=1.0, left=0.25 if kind == 'sub' else 0)
                    add(q, title, 12 if kind != 'sub' else 11)
                q = tr.cells[2].paragraphs[0]
                fmt(q, WD_ALIGN_PARAGRAPH.RIGHT, before=2 if kind != 'chapter' else 6, after=2, line=1.0)
                if bm:
                    field(q, 'PAGEREF %s \\h' % bm, '1', 12, bold)
                for c in tr.cells:
                    W.cell_borders(c)
                    W.cell_margins(c, t=0, b=0, l=40, r=40)
            ph._p.addnext(t._tbl)
        table_at(self.placeholders['toc'], self.toc, [0.7, 4.55, 0.7])
        table_at(self.placeholders['lot'], [('item', n, c, b) for n, c, b in self.lot], [1.0, 4.25, 0.7])
        table_at(self.placeholders['lof'], [('item', n, c, b) for n, c, b in self.lof], [1.0, 4.25, 0.7])

    # ------------------------------------------------------------------ body
    def start_main(self):
        self.new_section('decimal', 1)

    def chapter(self, title, label, kind):
        """kind: 'chapter' (numbered), 'appendix', 'plain' (References)."""
        if not self.first:
            self.page_break()
        self.first = False
        self.sec = self.sub = 0
        self.cnt = {}
        if kind == 'chapter':
            self.chap = str(sum(1 for e in self.toc if e[0] == 'chapter' and e[1].isdigit()) + 1)
            lead = 'Chapter %s' % self.chap
        elif kind == 'appendix':
            self.chap = chr(65 + sum(1 for e in self.toc if e[0] == 'appendix_marker'))
            lead = 'Appendix %s' % self.chap
        else:
            lead = None
        if lead:
            p = self.doc.add_paragraph()
            fmt(p, WD_ALIGN_PARAGRAPH.CENTER, after=6, keep=True, line=1.0)
            self.take_break(p)
            add(p, lead, 16, True)
        h = self.doc.add_paragraph(style='Heading 1')
        fmt(h, WD_ALIGN_PARAGRAPH.CENTER, after=18, keep=True, line=1.0)
        self.take_break(h)
        add(h, title, 16, True)
        bm = self.newmark()
        bookmark(h, bm)
        if kind == 'chapter':
            self.toc.append(('chapter', self.chap, 'CHAPTER %s: %s' % (self.chap, title.upper()), bm))
        elif kind == 'appendix':
            if not any(e[0] == 'end' and e[2] == 'APPENDICES' for e in self.toc):
                self.toc.append(('end', '', 'APPENDICES', None))
            self.toc.append(('appendix_marker', '', '', None))
            self.toc.append(('section', self.chap, title, bm))
        else:
            self.toc.append(('end', '', title.upper(), bm))

    def section(self, title, level):
        if level == 2:
            self.sec += 1
            self.sub = 0
            num = '%s.%d' % (self.chap, self.sec)
            size, before = 13, 12
        else:
            self.sub += 1
            num = '%s.%d.%d' % (self.chap, self.sec, self.sub)
            size, before = 12, 8
        h = self.doc.add_paragraph(style='Heading %d' % level)
        fmt(h, WD_ALIGN_PARAGRAPH.LEFT, before=before, after=4, keep=True, line=1.15)
        h.paragraph_format.left_indent = Inches(0)
        h.paragraph_format.first_line_indent = Inches(0)
        add(h, '%s %s' % (num, title), size, True)
        bm = self.newmark()
        bookmark(h, bm)
        self.toc.append(('section' if level == 2 else 'sub', num, title, bm))

    def runin(self, title):
        p = self.doc.add_paragraph()
        fmt(p, WD_ALIGN_PARAGRAPH.LEFT, before=6, after=2, keep=True, line=1.15)
        add(p, title, BODY, True, True)

    def paragraph(self, text):
        p = self.doc.add_paragraph()
        fmt(p, after=6, line=LINE)
        add(p, expand(text, self.L))

    def item(self, text, marker, level=0):
        p = self.doc.add_paragraph()
        ind = 0.3 + 0.3 * level
        fmt(p, WD_ALIGN_PARAGRAPH.LEFT if 'http' in text else WD_ALIGN_PARAGRAPH.JUSTIFY, after=3, left=ind, first=-0.3, line=LINE)
        p.paragraph_format.tab_stops.add_tab_stop(Inches(ind))
        r = p.add_run(marker + '\t')
        font(r)
        add(p, expand(text, self.L))

    def number(self, kind):
        self.cnt[kind] = self.cnt.get(kind, 0) + 1
        return '%s.%d' % (self.chap, self.cnt[kind])

    def figure(self, label, file, caption, width):
        num = self.number('fig')
        assert self.L.map['fig:' + label] == num, (label, num)
        p = self.doc.add_paragraph()
        fmt(p, WD_ALIGN_PARAGRAPH.CENTER, before=6, after=4, keep=True, line=1.0)
        p.add_run().add_picture(str(find_fig(file)), width=Inches(min(width or 5.7, TEXTW)))
        c = self.doc.add_paragraph()
        cap = expand(caption, self.L)
        fmt(c, WD_ALIGN_PARAGRAPH.CENTER if len(cap) < 110 else WD_ALIGN_PARAGRAPH.JUSTIFY, after=12, line=1.15)
        r = c.add_run('Figure %s: ' % num)
        font(r, 11, True)
        add(c, cap, 11)
        bm = self.newmark()
        bookmark(c, bm)
        self.lof.append(('Figure %s' % num, short(cap), bm))

    def table(self, label, caption, widths, header, rows, align=None, bold_cells=(), size=10):
        num = self.number('tab')
        assert self.L.map['tab:' + label] == num, (label, num)
        c = self.doc.add_paragraph()
        cap = expand(caption, self.L)
        fmt(c, WD_ALIGN_PARAGRAPH.CENTER if len(cap) < 110 else WD_ALIGN_PARAGRAPH.JUSTIFY, before=8, after=4, keep=True, line=1.15)
        r = c.add_run('Table %s: ' % num)
        font(r, 11, True)
        add(c, cap, 11)
        bm = self.newmark()
        bookmark(c, bm)
        self.lot.append(('Table %s' % num, short(cap), bm))
        widths = widths or [1] * len(header)
        widths = [TEXTW * w / sum(widths) for w in widths]
        align = align or ['l'] + ['r'] * (len(header) - 1)
        AL = {'l': WD_ALIGN_PARAGRAPH.LEFT, 'r': WD_ALIGN_PARAGRAPH.RIGHT, 'c': WD_ALIGN_PARAGRAPH.CENTER}
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        W.table_fixed(t, widths)
        keep = len(rows) <= 8
        for ri, row in enumerate([header] + rows):
            tr = t.rows[ri]
            W.row_props(tr, header=ri == 0)
            for ci, val in enumerate(row):
                cell = tr.cells[ci]
                W.cell_margins(cell, t=30, b=30, l=70, r=70)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                W.cell_borders(cell, top=4, bottom=4, left=4, right=4)
                if ri == 0:
                    shade(cell, 'D9E2F3')
                q = cell.paragraphs[0]
                fmt(q, AL[align[ci]] if ri else WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.0, keep=keep or ri == 0)
                add(q, expand(str(val), self.L), size, bold=(ri == 0) or ((ri - 1, ci) in bold_cells))
        sp = self.doc.add_paragraph()
        fmt(sp, after=6, line=1.0)

    def results(self, label, rq, metrics, caption, methods=None):
        R = BD.RES[rq]
        ms = methods or R['methods']
        P = set(R.get('pareto') or [])
        short_ = {'answer_f1': 'F1', 'answer_em': 'EM', 'answerable_f1': 'F1 (ans.)', 'mrr': 'MRR', 'gold_coverage': 'GC',
                  'source_f1': 'SF1', 'unnecessary_retrieval': 'WR', 'stopping_accuracy': 'Stop', 'required_source_miss': 'Miss',
                  'input_tokens': 'In tok.', 'output_tokens': 'Out tok.', 'total_tokens': 'Tokens', 'llm_calls': 'Calls',
                  'retrieval_calls': 'Search', 'latency': 'Lat. (s)'}
        header = ['Method'] + ['%s %s' % (short_.get(k, k), '↑' if k in BD.HIGHER else '↓') for k in metrics]
        rows, bold = [], set()
        for j, k in enumerate(metrics):
            vals = [R['descriptive'][m].get(k) for m in ms]
            good = [v for v in vals if v is not None]
            if good and max(good) - min(good) > 1e-12:
                best = max(good) if k in BD.HIGHER else min(good)
                bold.update((i, j + 1) for i, v in enumerate(vals) if v is not None and abs(v - best) < 1e-12)
        for m in ms:
            rows.append([m + (' ★' if m in P else '')] +
                        [BD.fmt(R['descriptive'][m][k], BD.KIND.get(k, 'rate')).replace(' s', '') for k in metrics])
        self.table(label, caption, [1.25] + [1.0] * len(metrics), header, rows, bold_cells=bold, size=9.5 if len(metrics) > 6 else 10)

    def equation(self, label, umath):
        num = self.number('eq')
        p = self.doc.add_paragraph(style='IEEE Equation')
        p.add_run(umath.strip() + '#(%s)' % num)

    def algorithm(self, label, title, lines):
        num = self.number('alg')
        c = self.doc.add_paragraph()
        fmt(c, WD_ALIGN_PARAGRAPH.LEFT, before=8, after=3, keep=True, line=1.15)
        r = c.add_run('Algorithm %s: ' % num)
        font(r, 11, True)
        add(c, expand(title, self.L), 11)
        t = self.doc.add_table(rows=1, cols=1)
        W.table_fixed(t, [TEXTW])
        cell = t.rows[0].cells[0]
        W.cell_borders(cell, top=8, bottom=8, left=4, right=4)
        W.cell_margins(cell, t=80, b=80, l=120, r=120)
        shade(cell, 'F7F9FC')
        first, step = True, 0
        for raw in lines:
            depth = len(raw) - len(raw.lstrip('>'))
            text = raw.lstrip('>').strip()
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            head, _, rest = text.partition(':')
            if head.lower() in ('input', 'output'):
                fmt(p, WD_ALIGN_PARAGRAPH.LEFT, after=1, left=0.7, first=-0.7, line=1.1, keep=True)
                add(p, head + ':', 10.5, True)
                add(p, ' ' + expand(rest.strip(), self.L), 10.5)
            else:
                step += 1
                ind = 0.55 + 0.25 * depth
                fmt(p, WD_ALIGN_PARAGRAPH.LEFT, after=1, left=ind, first=-0.55 if not depth else -0.25, line=1.1, keep=True)
                p.paragraph_format.tab_stops.add_tab_stop(Inches(ind))
                r = p.add_run(('Step %d.' % step if not depth else '•') + '\t')
                font(r, 10.5, not depth)
                add(p, expand(text, self.L), 10.5)
        sp = self.doc.add_paragraph()
        fmt(sp, after=6, line=1.0)

    def code(self, lines):
        t = self.doc.add_table(rows=1, cols=1)
        W.table_fixed(t, [TEXTW])
        cell = t.rows[0].cells[0]
        W.cell_borders(cell, top=4, bottom=4, left=4, right=4)
        W.cell_margins(cell, t=60, b=60, l=100, r=100)
        shade(cell, 'F4F4F4')
        for i, line in enumerate(lines):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            fmt(p, WD_ALIGN_PARAGRAPH.LEFT, after=0, line=1.0)
            r = p.add_run(line if line else ' ')
            font(r, 9, face=MONO)
        sp = self.doc.add_paragraph()
        fmt(sp, after=6, line=1.0)

    def references(self):
        for i, k in enumerate(self.L.cites, 1):
            p = self.doc.add_paragraph()
            fmt(p, WD_ALIGN_PARAGRAPH.LEFT, after=4, left=0.45, first=-0.45, line=1.15)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.45))
            r = p.add_run('[%d]\t' % i)
            font(r, 11)
            add(p, W.smart(REFS[k]), 11)


# =============================================================================== parser
def blocks(lines):
    out, buf = [], []

    def flush():
        if buf:
            out.append(('p', ' '.join(x.strip() for x in buf)))
            buf.clear()
    for line in lines:
        b = re.match(r'^(\s*)- (.*)$', line)
        n = re.match(r'^(\s*)(\d+)\. (.*)$', line)
        r = re.match(r'^(\s*)(\((?:[ivx]+|[a-h]|[A-Z]\d*|\d+)\)|[A-Z]{1,3}\d{1,2}[.:]?) (.*)$', line)
        if not line.strip():
            flush()
        elif b:
            flush()
            out.append(('i', b.group(2), len(b.group(1)) // 2, '•'))
        elif n:
            flush()
            out.append(('i', n.group(3), len(n.group(1)) // 2, n.group(2) + '.'))
        elif r and r.group(2).startswith('('):
            flush()
            out.append(('i', r.group(3), len(r.group(1)) // 2, r.group(2)))
        elif out and out[-1][0] == 'i' and line.startswith('  ') and not buf:
            o = out[-1]
            out[-1] = (o[0], o[1] + ' ' + line.strip(), o[2], o[3])
        else:
            buf.append(line)
    flush()
    return out


def body_until(lines, i, end='@end'):
    j = i + 1
    body = []
    while lines[j].strip() != end:
        body.append(lines[j].rstrip())
        j += 1
    return body, j + 1


def build(src):
    lines = src.read_text(encoding='utf-8').splitlines()
    R = Report()
    prescan(lines, R.L)
    meta, ack, abstract, kw, abbrevs = {}, [], [], '', []
    i = 0
    while i < len(lines):                                 # front-matter blocks first
        s = lines[i].strip()
        if s.startswith('%'):
            k, _, v = s[1:].partition(' ')
            meta[k] = v.strip()
        elif s == '@acknowledgement':
            body, i = body_until(lines, i)
            ack = [b for b in ' '.join(x if x else '\n' for x in body).split('\n') if b.strip()]
            continue
        elif s == '@abstract':
            body, i = body_until(lines, i)
            abstract = [b for b in ' '.join(x if x else '\n' for x in body).split('\n') if b.strip()]
            continue
        elif s.startswith('@keywords'):
            kw = s[len('@keywords'):].strip()
        elif s == '@abbrev':
            body, i = body_until(lines, i)
            abbrevs = sorted([tuple(x.strip() for x in b.split('|', 1)) for b in body if b.strip()], key=lambda t: t[0].upper())
            continue
        elif s.startswith('#'):
            break
        i += 1
    R.cover(meta)
    R.certificate(meta)
    R.declaration(meta)
    R.text_page('ACKNOWLEDGEMENT', ack)
    R.abstract(abstract, kw)
    R.lists(abbrevs)
    R.start_main()
    para = []

    def flush():
        for b in blocks(para):
            if b[0] == 'p':
                R.paragraph(b[1])
            else:
                R.item(b[1], b[3], b[2])
        para.clear()

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if s.startswith('<!--'):
            while '-->' not in lines[i]:
                i += 1
            i += 1
            continue
        if s.startswith('#'):
            flush()
            m = re.match(r'^(#{1,4}|#A)\s+(\*?)(.*?)\s*(\{#(?:ch|sec|app):[\w\-]+\})?$', s)
            mark, star, title = m.group(1), m.group(2), m.group(3)
            if mark == '#A':
                R.chapter(title, None, 'appendix')
            elif mark == '#':
                R.chapter(title, None, 'plain' if star else 'chapter')
            elif mark == '####':
                R.runin(expand(title, R.L))
            else:
                R.section(expand(title, R.L), len(mark))
            i += 1
            continue
        if s.startswith('@'):
            flush()
            cmd, _, rest = s.partition(' ')
            args = [a.strip() for a in rest.split('|')]
            if cmd == '@alg':
                body, i = body_until(lines, i)
                R.algorithm(args[0], args[1], [b.strip() for b in body if b.strip()])
                continue
            if cmd == '@code':
                body, i = body_until(lines, i)
                R.code(body)
                continue
            if cmd == '@table':
                widths = [float(x) for x in args[2].split(',')] if len(args) > 2 and args[2] else None
                j = i + 1
                rows = []
                while j < len(lines) and lines[j].strip().startswith('|'):
                    raw = lines[j].strip()[1:]
                    rows.append(W.split_row(raw[:-1] if raw.endswith('|') else raw))
                    j += 1
                header, sep, body = rows[0], rows[1], rows[2:]
                align = ['r' if c.endswith(':') and not c.startswith(':') else ('c' if c.startswith(':') and c.endswith(':') else 'l') for c in sep]
                R.table(args[0], args[1], widths, header, body, align=align)
                i = j
                continue
            if cmd == '@fig':
                R.figure(args[0], args[1], args[2], float(args[3]) if len(args) > 3 and args[3] else None)
            elif cmd == '@results':
                R.results(args[0], args[1], args[2].split(','), args[3], args[4].split(',') if len(args) > 4 else None)
            elif cmd == '@eq':
                R.equation(args[0], rest.split('|', 1)[1])
            elif cmd == '@refs':
                R.references()
            elif cmd == '@pagebreak':
                R.page_break()
            elif cmd in ('@abstract', '@acknowledgement', '@abbrev'):
                _, i = body_until(lines, i)
                continue
            elif cmd == '@keywords':
                pass
            else:
                raise ValueError('unknown directive ' + cmd)
            i += 1
            continue
        if not s.startswith('%'):
            para.append(line)
        i += 1
    flush()
    R.toc = [e for e in R.toc if e[0] != 'appendix_marker']
    R.fill_lists()
    unused = [k for k in REFS if k not in R.L.cites]
    R.doc.core_properties.title = meta.get('title', '')
    R.doc.core_properties.author = meta.get('student', '')
    R.doc.save(OUT)
    print('saved %s | chapters %d, figures %d, tables %d, references %d%s'
          % (OUT.name, sum(1 for e in R.toc if e[0] == 'chapter'), len(R.lof), len(R.lot), len(R.L.cites),
             (' | unused references: ' + ', '.join(unused)) if unused else ''))


if __name__ == '__main__':
    build(HERE / 'report.md')
